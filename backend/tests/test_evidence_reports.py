"""Tests for evidence-based reporting: normalized fields, exports, insufficient-evidence path."""
import os
import base64
from io import BytesIO
import requests
from PIL import Image
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

BASE_URL = os.environ["TEST_BASE_URL"].rstrip("/")

STATUS_VALUES = {"Sufficient evidence", "Insufficient evidence"}
CLASSIFICATION_VALUES = {"No apparent structural damage", "Minor / localized", "Moderate", "Severe", "Potentially unstable", "Insufficient evidence"}
CONFIDENCE_VALUES = {"Low", "Moderate", "High"}
URGENCY_VALUES = {"Critical", "High", "Moderate", "Monitor"}


def _make_image_b64(uniform=False):
    img = Image.new("RGB", (48, 48), "white")
    if not uniform:
        for x in range(10, 30):
            for y in range(12, 34):
                img.putpixel((x, y), (210, 40, 40))
    buf = BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode()


DETAILED_PAYLOAD = {
    "structure_name": "TEST_evidence_detailed",
    "component_type": "Reinforced concrete column",
    "construction_material": "Reinforced concrete",
    "blast_context": "Vehicle-borne blast ~8 m north, ground level, ~30 min prior.",
    "crack_width": "1–3 mm",
    "spalling_depth": "25–50 mm",
    "reinforcement_exposure": "Visible, no apparent distortion",
    "integrity_status": "Needs assessment",
    "site_notes": "Diagonal cracks near base; concrete cover missing in patches; adjacent slab appears intact.",
}

MINIMAL_PAYLOAD = {
    "structure_name": "TEST_evidence_minimal",
    "component_type": "Load-bearing wall",
    "construction_material": "Unknown",
    "blast_context": "Unclear",
    "crack_width": "Not measured",
    "spalling_depth": "None observed",
    "reinforcement_exposure": "Not visible",
    "integrity_status": "Needs assessment",
    "site_notes": "",
}


# ---------- Normalization / new fields ----------

def _validate_common(rec):
    assert rec["assessment_status"] in STATUS_VALUES
    assert rec["damage_classification"] in CLASSIFICATION_VALUES

    q = rec["quick_assessment"]
    for k in ("element", "apparent_condition", "severity", "confidence", "immediate_concern", "engineer_verification_required"):
        assert isinstance(q[k], str) and q[k]
    assert q["confidence"] in CONFIDENCE_VALUES
    assert isinstance(q["key_visible_evidence"], list)

    c = rec["evidence_confidence"]
    assert c["overall"] in CONFIDENCE_VALUES
    assert isinstance(c["percentage"], int) and 0 <= c["percentage"] <= 100
    for k in ("evidence_detected", "estimated_inferred", "cannot_be_determined"):
        assert isinstance(c[k], list)
    assert c["cannot_be_determined"], "cannot_be_determined must have fallback content"

    m = rec["possible_damage_mechanism"]
    assert isinstance(m["most_likely"], str) and m["most_likely"]
    assert isinstance(m["alternatives"], list)
    assert m["confidence"] in CONFIDENCE_VALUES
    assert isinstance(m["supporting_evidence"], list)

    assert isinstance(rec["limitations"], list) and rec["limitations"]
    assert isinstance(rec["required_verification"], list) and rec["required_verification"]

    # Exactly 3 recommendations
    assert len(rec["recommendations"]) == 3
    for idx, r in enumerate(rec["recommendations"], 1):
        assert r["rank"] == idx
        assert r["urgency"] in URGENCY_VALUES
        for k in ("technique", "fit", "reasoning", "evidence_link", "appropriate_when", "do_not_use_when", "timeframe"):
            assert isinstance(r[k], str) and r[k]
        for k in ("materials", "procedure", "safety_notes", "is_code_considerations"):
            assert isinstance(r[k], list)


def test_detailed_assessment_populates_all_new_fields():
    r = requests.post(f"{BASE_URL}/api/assessments", json=DETAILED_PAYLOAD, timeout=180)
    assert r.status_code == 200, r.text
    rec = r.json()
    _validate_common(rec)
    # Persistence
    g = requests.get(f"{BASE_URL}/api/assessments", timeout=30).json()
    assert any(x["id"] == rec["id"] for x in g)
    return rec


def test_minimal_input_preserves_insufficient_evidence_language():
    r = requests.post(f"{BASE_URL}/api/assessments", json=MINIMAL_PAYLOAD, timeout=180)
    assert r.status_code == 200, r.text
    rec = r.json()
    _validate_common(rec)
    # Conservative: minimal input should yield Low/Moderate confidence
    assert rec["evidence_confidence"]["overall"] in {"Low", "Moderate"}


# ---------- Exports ----------

def _fetch_report(assessment_id, fmt):
    r = requests.get(f"{BASE_URL}/api/assessments/{assessment_id}/report/{fmt}", timeout=60)
    assert r.status_code == 200, r.text
    return r.content


def test_pdf_report_contains_new_sections():
    r = requests.post(f"{BASE_URL}/api/assessments", json={**DETAILED_PAYLOAD, "structure_name": "TEST_pdf_report"}, timeout=180)
    assert r.status_code == 200
    aid = r.json()["id"]
    pdf_bytes = _fetch_report(aid, "pdf")
    assert pdf_bytes.startswith(b"%PDF")
    reader = PdfReader(BytesIO(pdf_bytes))
    text = "\n".join((p.extract_text() or "") for p in reader.pages).upper()
    for section in ("QUICK ASSESSMENT", "EVIDENCE & CONFIDENCE", "POSSIBLE DAMAGE MECHANISM", "LIMITATIONS", "REQUIRED VERIFICATION"):
        assert section in text, f"Missing section '{section}' in PDF"


def test_docx_report_contains_new_sections():
    r = requests.post(f"{BASE_URL}/api/assessments", json={**DETAILED_PAYLOAD, "structure_name": "TEST_docx_report"}, timeout=180)
    aid = r.json()["id"]
    docx_bytes = _fetch_report(aid, "docx")
    doc = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in doc.paragraphs).lower()
    for kw in ("quick assessment", "evidence & confidence", "possible damage mechanism", "limitations", "required verification"):
        assert kw in text, f"Missing '{kw}' in DOCX"


def test_xlsx_report_contains_new_sections():
    r = requests.post(f"{BASE_URL}/api/assessments", json={**DETAILED_PAYLOAD, "structure_name": "TEST_xlsx_report"}, timeout=180)
    aid = r.json()["id"]
    xlsx_bytes = _fetch_report(aid, "xlsx")
    wb = load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    joined = " ".join(cells).lower()
    for kw in ("assessment status", "damage classification", "confidence", "evidence detected", "estimated / inferred", "cannot be determined", "required verification", "limitations"):
        assert kw in joined, f"Missing '{kw}' in XLSX"


# ---------- Image validation ----------

def test_rejects_malformed_base64():
    r = requests.post(f"{BASE_URL}/api/assessments", json={**DETAILED_PAYLOAD, "structure_name": "TEST_bad_b64", "image_base64": "!!!not-base64!!!", "image_mime": "image/png"}, timeout=30)
    assert r.status_code == 422


def test_rejects_wrong_mime():
    r = requests.post(f"{BASE_URL}/api/assessments", json={**DETAILED_PAYLOAD, "structure_name": "TEST_bad_mime", "image_base64": _make_image_b64(), "image_mime": "image/gif"}, timeout=30)
    assert r.status_code == 422


def test_rejects_uniform_image():
    r = requests.post(f"{BASE_URL}/api/assessments", json={**DETAILED_PAYLOAD, "structure_name": "TEST_uniform_img", "image_base64": _make_image_b64(uniform=True), "image_mime": "image/png"}, timeout=30)
    assert r.status_code == 422
