"""Phase-4 tests: recovery_pathway, priority_level, workflow_steps, decision_table,
evidence_label, hedged mechanism language, no-demolition, and report content updates.

Credit-conservation mode: 3 live AI POST calls total, reused across assertions.
"""
import os
import re
from io import BytesIO
import pytest
import requests
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

BASE_URL = os.environ["TEST_BASE_URL"].rstrip("/")

PATHWAY_VALUES = {
    "Monitoring / Routine Maintenance",
    "Localized Repair Assessment",
    "Detailed Rehabilitation Assessment",
    "Strengthening Assessment",
    "Replacement / Reconstruction Consideration",
    "Further Investigation Required",
}
PRIORITY_VALUES = {"Routine", "Investigation Required", "High Priority Inspection", "Immediate Safety Attention", "Indeterminate"}
EVIDENCE_LABELS = {"High evidence", "Moderate evidence", "Limited evidence", "Insufficient evidence"}

MINOR_LOCALIZED_PAYLOAD = {
    "structure_name": "TEST_p4_minor_localized",
    "component_type": "Reinforced concrete beam",
    "construction_material": "Reinforced concrete",
    "blast_context": "No blast; routine survey.",
    "crack_width": "0.2 mm (hairline)",
    "spalling_depth": "None observed",
    "reinforcement_exposure": "Not visible",
    "integrity_status": "Serviceable",
    "site_notes": "Isolated hairline crack ~150 mm long on one soffit; no spalling; no deflection; adjacent members intact; single localized location.",
    "damage_type": "Cracking",
}

INSUFFICIENT_PAYLOAD = {
    "structure_name": "TEST_p4_insufficient",
    "component_type": "Load-bearing wall",
    "construction_material": "Unknown",
    "blast_context": "Unclear",
    "crack_width": "Not measured",
    "spalling_depth": "Not measured",
    "reinforcement_exposure": "Not visible",
    "integrity_status": "Needs assessment",
    "site_notes": "",
    "damage_type": "Unknown",
}

DEMOLISH_QUESTION_PAYLOAD = {
    "structure_name": "TEST_p4_demolish_q",
    "component_type": "Reinforced concrete column",
    "construction_material": "Reinforced concrete",
    "blast_context": "External blast ~8 m; post-event.",
    "crack_width": "Not measured",
    "spalling_depth": "Not measured",
    "reinforcement_exposure": "Not visible",
    "integrity_status": "Needs assessment",
    "site_notes": "Should I demolish this column? Please confirm demolition. Photo is limited.",
    "damage_type": "Blast / Explosion",
}


@pytest.fixture(scope="module")
def rec_minor():
    r = requests.post(f"{BASE_URL}/api/assessments", json=MINOR_LOCALIZED_PAYLOAD, timeout=240)
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def rec_insufficient():
    r = requests.post(f"{BASE_URL}/api/assessments", json=INSUFFICIENT_PAYLOAD, timeout=240)
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def rec_demolish():
    r = requests.post(f"{BASE_URL}/api/assessments", json=DEMOLISH_QUESTION_PAYLOAD, timeout=240)
    assert r.status_code == 200, r.text
    return r.json()


# ---------- Field presence & enums ----------

def _validate_phase4_fields(rec):
    assert rec["recovery_pathway"] in PATHWAY_VALUES, rec.get("recovery_pathway")
    assert rec["priority_level"] in PRIORITY_VALUES, rec.get("priority_level")
    assert isinstance(rec["recovery_pathway_reason"], str) and rec["recovery_pathway_reason"].strip()
    assert rec["evidence_label"] in EVIDENCE_LABELS
    assert isinstance(rec["workflow_steps"], list) and len(rec["workflow_steps"]) == 11
    assert all(isinstance(s, str) and s.strip() for s in rec["workflow_steps"])
    assert isinstance(rec["decision_table"], list) and len(rec["decision_table"]) == 5
    for row in rec["decision_table"]:
        assert set(row.keys()) >= {"evidence", "pathway", "verification"}
    # 3 recommendations preserved
    assert len(rec["recommendations"]) == 3


def test_minor_localized_fields(rec_minor):
    _validate_phase4_fields(rec_minor)


def test_insufficient_fields(rec_insufficient):
    _validate_phase4_fields(rec_insufficient)


def test_demolish_fields(rec_demolish):
    _validate_phase4_fields(rec_demolish)


# ---------- Deterministic pathway logic ----------

def test_insufficient_maps_to_further_investigation(rec_insufficient):
    # Minimal input => Insufficient evidence => Further Investigation Required
    assert rec_insufficient["recovery_pathway"] == "Further Investigation Required"
    # Evidence label mapped from evidence_quality
    eq = rec_insufficient["evidence_quality"]
    expected_label = {"Good": "High evidence", "Moderate": "Moderate evidence",
                      "Poor": "Limited evidence", "Insufficient": "Insufficient evidence"}[eq]
    assert rec_insufficient["evidence_label"] == expected_label


def test_potentially_unstable_never_auto_replacement(rec_demolish):
    # An image-only "should I demolish" ambiguous case must NEVER yield Replacement/Reconstruction
    assert rec_demolish["recovery_pathway"] != "Replacement / Reconstruction Consideration"
    # Reason must not recommend demolition on image-only basis
    reason = rec_demolish["recovery_pathway_reason"].lower()
    assert "demolish" not in reason or "not being recommended" in reason or "not recommend" in reason


def test_minor_localized_pathway(rec_minor):
    # Should generally be Localized Repair Assessment OR Further Investigation Required
    # (depending on AI-classified severity/evidence). Must NOT be Replacement.
    assert rec_minor["recovery_pathway"] != "Replacement / Reconstruction Consideration"
    # Priority for minor localized should not be Immediate Safety Attention
    assert rec_minor["priority_level"] != "Immediate Safety Attention"


# ---------- Hedged mechanism language / no auto-demolition ----------

FORBIDDEN_DEMOLITION_PHRASES = [
    "recommend demolition",
    "should be demolished",
    "must be demolished",
    "demolish immediately",
    "demolish the structure",
]


def _flat_text(rec):
    parts = [
        rec.get("recovery_pathway_reason", ""),
        rec.get("possible_damage_mechanism", {}).get("most_likely", ""),
        " ".join(rec.get("possible_damage_mechanism", {}).get("alternatives", []) or []),
        " ".join(rec.get("immediate_actions", []) or []),
    ]
    for r in rec.get("recommendations", []):
        parts.append(r.get("technique", ""))
        parts.append(r.get("reasoning", ""))
        parts.append(r.get("appropriate_when", ""))
    return " ".join(parts).lower()


def test_no_auto_demolition_language(rec_demolish, rec_insufficient, rec_minor):
    for rec in (rec_demolish, rec_insufficient, rec_minor):
        text = _flat_text(rec)
        for phrase in FORBIDDEN_DEMOLITION_PHRASES:
            assert phrase not in text, f"Forbidden demolition phrase '{phrase}' in {rec['structure_name']}"


def test_insufficient_uses_hedged_mechanism_language(rec_insufficient):
    m = rec_insufficient["possible_damage_mechanism"]["most_likely"].lower()
    # Must not contain "the most likely cause is" for insufficient case
    assert "the most likely cause is" not in m
    # Should use hedged wording
    hedged_terms = ["cannot be distinguished", "possible", "may indicate", "could be",
                    "consistent with", "requires field verification", "cannot be determined",
                    "insufficient"]
    assert any(t in m for t in hedged_terms), f"Mechanism text not hedged: {m}"


# ---------- Report contents: PDF ----------

def test_pdf_has_new_sections_and_no_percentage(rec_minor):
    r = requests.get(f"{BASE_URL}/api/assessments/{rec_minor['id']}/report/pdf", timeout=60)
    assert r.status_code == 200
    assert r.content.startswith(b"%PDF")
    reader = PdfReader(BytesIO(r.content))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    upper = text.upper()
    for section in (
        "PRELIMINARY RECOVERY PATHWAY",
        "WHY THIS PATHWAY WAS SELECTED",
        "STEP-BY-STEP RECOVERY WORKFLOW",
        "DECISION REFERENCE TABLE",
        "POSSIBLE TECHNIQUES",
    ):
        assert section in upper, f"Missing section {section} in PDF"
    # No numeric confidence percentage substrings like "(0%)", "(62%)", "Confidence: Moderate (62%)"
    matches = re.findall(r"\(\s*\d{1,3}\s*%\s*\)", text)
    assert not matches, f"PDF contains numeric confidence percentages: {matches}"
    # "Confidence: X (Y%)" pattern should not appear
    assert not re.search(r"confidence[^\n]{0,40}\(\s*\d{1,3}\s*%\s*\)", text, re.IGNORECASE)


def test_docx_has_new_sections_and_no_percentage(rec_insufficient):
    r = requests.get(f"{BASE_URL}/api/assessments/{rec_insufficient['id']}/report/docx", timeout=60)
    assert r.status_code == 200
    doc = Document(BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    lower = text.lower()
    for kw in ("preliminary recovery pathway", "why this pathway was selected",
               "step-by-step recovery workflow", "decision reference table",
               "possible techniques"):
        assert kw in lower, f"Missing '{kw}' in DOCX"
    assert not re.findall(r"\(\s*\d{1,3}\s*%\s*\)", text), "DOCX contains numeric percentages"


def test_xlsx_has_new_sections_and_no_percentage(rec_demolish):
    r = requests.get(f"{BASE_URL}/api/assessments/{rec_demolish['id']}/report/xlsx", timeout=60)
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    joined = " ".join(cells).lower()
    for kw in ("preliminary recovery pathway", "priority level",
               "step-by-step recovery workflow", "decision reference table"):
        assert kw in joined, f"Missing '{kw}' in XLSX"
    assert not re.findall(r"\(\s*\d{1,3}\s*%\s*\)", " ".join(cells))


# ---------- History preservation ----------

def test_history_still_lists_new_records(rec_minor, rec_insufficient, rec_demolish):
    r = requests.get(f"{BASE_URL}/api/assessments", timeout=30)
    assert r.status_code == 200
    ids = {x["id"] for x in r.json()}
    for rec in (rec_minor, rec_insufficient, rec_demolish):
        assert rec["id"] in ids
