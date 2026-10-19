from fastapi import FastAPI, APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Any
import uuid
import json
import base64
import binascii
from io import BytesIO
from datetime import datetime, timezone
from PIL import Image, UnidentifiedImageError
from emergentintegrations.llm.chat import LlmChat, UserMessage, TextDelta
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.colors import Color
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
from docx import Document
from docx.shared import RGBColor, Pt
from openpyxl import Workbook


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app without a prefix
app = FastAPI()


@app.get("/")
async def root():
    return {
        "message": "Hello from Restruct-AI!",
        "description": "This is the backend for Restruct-AI APP.",
        "github": "https://github.com/joelgiradkar7/Restruct-AI/tree/main"
    }


# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")


# Define Models
class StatusCheck(BaseModel):
    model_config = ConfigDict(extra="ignore")  # Ignore MongoDB's _id field
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_name: str
    timestamp: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class StatusCheckCreate(BaseModel):
    client_name: str

class AssessmentCreate(BaseModel):
    structure_name: str
    component_type: str
    construction_material: str
    blast_context: str
    crack_width: str
    spalling_depth: str
    reinforcement_exposure: str
    integrity_status: str
    site_notes: str
    damage_type: str = "Unknown"
    image_base64: Optional[str] = None
    image_mime: Optional[str] = None

class Assessment(BaseModel):
    id: str
    structure_name: str
    component_type: str
    integrity_status: str
    created_at: datetime
    recommendations: List[dict]
    immediate_actions: List[str]
    disclaimer: str
    assessment_status: Optional[str] = None
    damage_classification: Optional[str] = None
    quick_assessment: Optional[dict] = None
    evidence_confidence: Optional[dict] = None
    possible_damage_mechanism: Optional[dict] = None
    limitations: Optional[List[str]] = None
    required_verification: Optional[List[str]] = None
    damage_type: Optional[str] = None
    construction_material: Optional[str] = None
    observed_conditions: Optional[List[str]] = None
    user_measurements: Optional[List[str]] = None
    preliminary_severity: Optional[str] = None
    evidence_quality: Optional[str] = None
    evidence_label: Optional[str] = None
    missing_information: Optional[List[str]] = None
    recovery_pathway: Optional[str] = None
    priority_level: Optional[str] = None
    recovery_pathway_reason: Optional[str] = None
    workflow_steps: Optional[List[str]] = None
    decision_table: Optional[List[dict]] = None

# Add your routes to the router instead of directly to app
@api_router.get("/")
async def root():
    return {"message": "ResructAI API ready"}

def clean_json_text(raw: str) -> dict:
    text = raw.strip().replace("```json", "").replace("```", "").strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start >= 0 and end > start:
            return json.loads(text[start:end + 1])
        raise

STATUS_VALUES = {"Sufficient evidence", "Insufficient evidence"}
CLASSIFICATION_VALUES = {"No apparent structural damage", "Minor / localized", "Moderate", "Severe", "Potentially unstable", "Insufficient evidence"}
CONFIDENCE_VALUES = {"Low", "Moderate", "High"}
URGENCY_VALUES = {"Critical", "High", "Moderate", "Monitor"}
SEVERITY_VALUES = {"Low", "Moderate", "High", "Critical", "Indeterminate"}
EVIDENCE_QUALITY_VALUES = {"Good", "Moderate", "Poor", "Insufficient"}
PATHWAY_VALUES = {
    "Monitoring / Routine Maintenance",
    "Localized Repair Assessment",
    "Detailed Rehabilitation Assessment",
    "Strengthening Assessment",
    "Replacement / Reconstruction Consideration",
    "Further Investigation Required",
}
PRIORITY_VALUES = {"Routine", "Investigation Required", "High Priority Inspection", "Immediate Safety Attention", "Indeterminate"}
EVIDENCE_LABEL_MAP = {"Good": "High evidence", "Moderate": "Moderate evidence", "Poor": "Limited evidence", "Insufficient": "Insufficient evidence"}
WORKFLOW_STEPS = [
    "Immediate precautions where appropriate (exclusion zones, shoring, halting operations).",
    "Detailed visual inspection by a qualified engineer on site.",
    "Damage mapping and dimensional measurements (crack widths/lengths, spalling extents, deflections).",
    "Determine whether damage is active/progressive (monitoring, telltales, repeat measurements).",
    "Investigate possible causes (loading history, environmental exposure, construction records).",
    "Assess structural and material condition (in-situ / laboratory testing as appropriate).",
    "Engineering analysis where required (capacity, load path, seismic, durability).",
    "Select preliminary recovery pathway based on findings.",
    "Develop engineer-approved repair / strengthening / rehabilitation design.",
    "Implement the approved intervention under qualified supervision.",
    "Post-intervention inspection and monitoring.",
]
DECISION_TABLE = [
    {"evidence": "Localized / apparently stable distress", "pathway": "Localized Repair Assessment", "verification": "Confirm cause and structural significance."},
    {"evidence": "Potential structural distress", "pathway": "Detailed Rehabilitation Assessment or Strengthening Assessment", "verification": "Detailed structural investigation."},
    {"evidence": "Extensive / complex deterioration", "pathway": "Detailed Rehabilitation Assessment", "verification": "Structural analysis and feasibility assessment."},
    {"evidence": "Insufficient evidence", "pathway": "Further Investigation Required", "verification": "Site inspection and material / structural testing."},
    {"evidence": "Potential severe instability indicators", "pathway": "Immediate Safety Attention + Further Investigation", "verification": "Qualified engineer / competent authority assessment."},
]

def _derive_pathway(classification: str, severity: str, quality: str) -> str:
    if quality == "Insufficient" or classification == "Insufficient evidence" or severity == "Indeterminate":
        return "Further Investigation Required"
    if classification == "Potentially unstable" or severity == "Critical":
        return "Further Investigation Required"
    if classification == "No apparent structural damage":
        return "Monitoring / Routine Maintenance"
    if classification == "Minor / localized":
        return "Localized Repair Assessment"
    if classification == "Severe":
        return "Detailed Rehabilitation Assessment"
    if classification == "Moderate":
        return "Detailed Rehabilitation Assessment" if severity in {"High", "Critical"} else "Localized Repair Assessment"
    return "Further Investigation Required"

def _derive_priority(severity: str, classification: str) -> str:
    if classification == "Potentially unstable" or severity == "Critical":
        return "Immediate Safety Attention"
    if severity == "High" or classification == "Severe":
        return "High Priority Inspection"
    if severity == "Moderate" or classification in {"Moderate", "Minor / localized"}:
        return "Investigation Required"
    if severity == "Low" or classification == "No apparent structural damage":
        return "Routine"
    return "Indeterminate"

def _pathway_reason(pathway: str, classification: str, quality: str, severity: str, missing: List[str]) -> str:
    if pathway == "Further Investigation Required":
        if classification == "Potentially unstable" or severity == "Critical":
            return "Potential severe-instability indicators are present but the photograph and provided information cannot establish load path, member capacity, hidden damage, or reinforcement condition. Site inspection by a qualified engineer is required before any recovery pathway is chosen."
        return "The available evidence is insufficient (evidence quality: " + quality + ") to select a specific recovery pathway. Additional site inspection and information — for example: " + ("; ".join(missing[:3]) if missing else "further site data") + " — are needed."
    if pathway == "Monitoring / Routine Maintenance":
        return "No apparent structural damage is visible in the available evidence; routine monitoring is appropriate while confirming with a qualified engineer."
    if pathway == "Localized Repair Assessment":
        return "Damage appears localized based on the available evidence with no clear indicators of major structural instability. Repair suitability requires site inspection and engineering verification."
    if pathway == "Detailed Rehabilitation Assessment":
        return "The observed distress appears significant or multi-factorial and may not be appropriately addressed through a simple localized repair. A detailed rehabilitation assessment by a qualified engineer is recommended."
    if pathway == "Strengthening Assessment":
        return "Evidence suggests potential structural performance concerns that may require strengthening. Remaining capacity, capacity loss, and required reinforcement must be established through proper engineering analysis — none of these values are computed by this preliminary review."
    return "Replacement or reconstruction consideration is a major decision that requires detailed structural assessment, safety analysis, feasibility, economics, regulatory review, and approval by the appropriate qualified professionals. It is not being recommended on the basis of image-only evidence."

def _pick(value: Any, allowed: set, fallback: str) -> str:
    return value if isinstance(value, str) and value in allowed else fallback

def _string_list(value: Any) -> List[str]:
    if isinstance(value, list):
        return [str(x).strip() for x in value if isinstance(x, (str, int, float)) and str(x).strip()]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []

def _int_percent(value: Any) -> int:
    try:
        return max(0, min(100, int(value)))
    except (TypeError, ValueError):
        return 0

def normalize_advice(advice: dict, component_type: str) -> dict:
    if not isinstance(advice, dict):
        advice = {}
    status = _pick(advice.get("assessment_status"), STATUS_VALUES, "Insufficient evidence")
    classification = _pick(advice.get("damage_classification"), CLASSIFICATION_VALUES, "Insufficient evidence")
    if status == "Insufficient evidence":
        classification = "Insufficient evidence"

    quick = advice.get("quick_assessment") if isinstance(advice.get("quick_assessment"), dict) else {}
    quick_norm = {
        "element": str(quick.get("element") or component_type),
        "apparent_condition": str(quick.get("apparent_condition") or "Cannot be determined from available evidence"),
        "key_visible_evidence": _string_list(quick.get("key_visible_evidence")),
        "severity": str(quick.get("severity") or classification),
        "confidence": _pick(quick.get("confidence"), CONFIDENCE_VALUES, "Low"),
        "immediate_concern": str(quick.get("immediate_concern") or "Requires engineer review before further action"),
        "engineer_verification_required": str(quick.get("engineer_verification_required") or "Yes — licensed structural engineer must inspect and verify"),
    }

    conf = advice.get("evidence_confidence") if isinstance(advice.get("evidence_confidence"), dict) else {}
    conf_norm = {
        "overall": _pick(conf.get("overall"), CONFIDENCE_VALUES, "Low"),
        "percentage": _int_percent(conf.get("percentage")),
        "evidence_detected": _string_list(conf.get("evidence_detected")),
        "estimated_inferred": _string_list(conf.get("estimated_inferred")),
        "cannot_be_determined": _string_list(conf.get("cannot_be_determined")) or ["Residual capacity and hidden damage cannot be determined from available evidence."],
    }

    mech = advice.get("possible_damage_mechanism") if isinstance(advice.get("possible_damage_mechanism"), dict) else {}
    mech_norm = {
        "most_likely": str(mech.get("most_likely") or "Cannot be determined from available evidence."),
        "alternatives": _string_list(mech.get("alternatives")),
        "confidence": _pick(mech.get("confidence"), CONFIDENCE_VALUES, "Low"),
        "supporting_evidence": _string_list(mech.get("supporting_evidence")),
    }

    limitations = _string_list(advice.get("limitations")) or ["Image-based assessment cannot establish residual load-bearing capacity, reinforcement loss, concrete strength, hidden delamination, or complete structural stability."]
    verification = _string_list(advice.get("required_verification")) or ["Detailed visual inspection and structural calculations by a licensed engineer."]
    immediate = _string_list(advice.get("immediate_actions")) or ["Restrict access to the affected zone until a licensed structural engineer inspects the component."]

    raw_recs = advice.get("recommendations") if isinstance(advice.get("recommendations"), list) else []
    recs: List[dict] = []
    for idx, item in enumerate(raw_recs[:3]):
        if not isinstance(item, dict):
            continue
        recs.append({
            "rank": idx + 1,
            "technique": str(item.get("technique") or "Engineer-led verification and interim protection"),
            "fit": str(item.get("fit") or "Requires engineer confirmation"),
            "reasoning": str(item.get("reasoning") or "Evidence is insufficient to commit to a specific rehabilitation technique."),
            "evidence_link": str(item.get("evidence_link") or "UNKNOWN — the engineer must confirm the condition that makes this option appropriate."),
            "appropriate_when": str(item.get("appropriate_when") or "Only after the listed verification confirms the observed condition."),
            "do_not_use_when": str(item.get("do_not_use_when") or "Do not use until instability, hidden damage, and load-path effects are checked."),
            "urgency": _pick(item.get("urgency"), URGENCY_VALUES, "Moderate"),
            "materials": _string_list(item.get("materials")),
            "procedure": _string_list(item.get("procedure")),
            "safety_notes": _string_list(item.get("safety_notes")),
            "timeframe": str(item.get("timeframe") or "To be confirmed after engineer review"),
            "is_code_considerations": _string_list(item.get("is_code_considerations")) or ["Verify the current applicable Indian Standards and project provisions with the engineer."],
        })
    while len(recs) < 3:
        rank = len(recs) + 1
        recs.append({
            "rank": rank,
            "technique": "Engineer-led verification pathway",
            "fit": "Requires engineer confirmation",
            "reasoning": "Insufficient evidence to recommend a specific rehabilitation technique for this pathway.",
            "evidence_link": "UNKNOWN — evidence does not yet support a distinct technique for this pathway.",
            "appropriate_when": "Only after site inspection, material testing, and structural analysis by a licensed engineer.",
            "do_not_use_when": "Do not proceed until instability, hidden damage, and load paths are verified.",
            "urgency": "Monitor",
            "materials": [],
            "procedure": ["Await engineer inspection and testing before selecting a repair technique."],
            "safety_notes": ["Maintain exclusion zone and shoring until verification is complete."],
            "timeframe": "To be confirmed after engineer review",
            "is_code_considerations": ["Verify the current applicable Indian Standards and project provisions with the engineer."],
        })

    missing_information_list = _string_list(advice.get("missing_information")) or ["Additional site photographs, dimensions, structural drawings, and material information would improve the assessment."]
    preliminary_severity = _pick(advice.get("preliminary_severity"), SEVERITY_VALUES, "Indeterminate")
    evidence_quality = _pick(advice.get("evidence_quality"), EVIDENCE_QUALITY_VALUES, "Insufficient")
    pathway = _pick(advice.get("recovery_pathway"), PATHWAY_VALUES, "") or _derive_pathway(classification, preliminary_severity, evidence_quality)
    priority = _pick(advice.get("priority_level"), PRIORITY_VALUES, "") or _derive_priority(preliminary_severity, classification)
    pathway_reason = str(advice.get("recovery_pathway_reason") or "").strip() or _pathway_reason(pathway, classification, evidence_quality, preliminary_severity, missing_information_list)
    return {
        "assessment_status": status,
        "damage_classification": classification,
        "quick_assessment": quick_norm,
        "evidence_confidence": conf_norm,
        "possible_damage_mechanism": mech_norm,
        "limitations": limitations,
        "required_verification": verification,
        "immediate_actions": immediate,
        "recommendations": recs,
        "observed_conditions": _string_list(advice.get("observed_conditions")),
        "user_measurements": _string_list(advice.get("user_measurements")),
        "preliminary_severity": preliminary_severity,
        "evidence_quality": evidence_quality,
        "evidence_label": EVIDENCE_LABEL_MAP.get(evidence_quality, "Insufficient evidence"),
        "missing_information": missing_information_list,
        "recovery_pathway": pathway,
        "priority_level": priority,
        "recovery_pathway_reason": pathway_reason,
        "workflow_steps": WORKFLOW_STEPS,
        "decision_table": DECISION_TABLE,
    }

async def generate_advice(payload: AssessmentCreate) -> dict:
    key = os.environ.get("EMERGENT_LLM_KEY")
    if not key:
        raise HTTPException(status_code=503, detail="AI service is not configured")
    if payload.image_base64:
        if payload.image_mime not in {"image/jpeg", "image/png", "image/webp"}:
            raise HTTPException(status_code=422, detail="Only JPG, PNG, or WEBP photos are supported")
        try:
            raw_image = base64.b64decode(payload.image_base64.split(",")[-1], validate=True)
            with Image.open(BytesIO(raw_image)) as image:
                image.verify()
            with Image.open(BytesIO(raw_image)) as image:
                sample = image.convert("L").resize((64, 64)).getextrema()
                if sample[1] - sample[0] <= 2:
                    raise HTTPException(status_code=422, detail="Please upload a photo with visible structural detail")
        except HTTPException:
            raise
        except (ValueError, binascii.Error, UnidentifiedImageError, OSError) as exc:
            raise HTTPException(status_code=422, detail="The uploaded photo is not a valid image") from exc
    prompt = f"""You are a cautious structural rehabilitation advisor supporting work in India. Analyze the supplied user fields and optional image for a qualified engineer across ANY type of structural damage — including but not limited to cracking, spalling, delamination, reinforcement corrosion, fire/thermal, earthquake/seismic, blast/explosion, vehicle/impact, water/moisture, chemical deterioration, differential settlement, foundation distress, overloading, construction defects, material deterioration, structural deformation, masonry damage, connection/joint damage, or unknown/other damage. Never claim a final design, occupancy approval, structural safety guarantee, remaining capacity, or confirmed cause. Return ONLY valid JSON with this exact shape:
{{"assessment_status":"Sufficient evidence|Insufficient evidence","damage_classification":"No apparent structural damage|Minor / localized|Moderate|Severe|Potentially unstable|Insufficient evidence","observed_conditions":["..."],"user_measurements":["..."],"preliminary_severity":"Low|Moderate|High|Critical|Indeterminate","evidence_quality":"Good|Moderate|Poor|Insufficient","missing_information":["..."],"quick_assessment":{{"element":"...","apparent_condition":"...","key_visible_evidence":["..."],"severity":"...","confidence":"Low|Moderate|High","immediate_concern":"...","engineer_verification_required":"..."}},"evidence_confidence":{{"overall":"Low|Moderate|High","percentage":0,"evidence_detected":["..."],"estimated_inferred":["..."],"cannot_be_determined":["..."]}},"possible_damage_mechanism":{{"most_likely":"...","alternatives":["..."],"confidence":"Low|Moderate|High","supporting_evidence":["..."]}},"limitations":["..."],"required_verification":["..."],"immediate_actions":["..."],"recommendations":[{{"rank":1,"technique":"...","fit":"...","reasoning":"...","evidence_link":"...","appropriate_when":"...","do_not_use_when":"...","urgency":"Critical|High|Moderate|Monitor","materials":["..."],"procedure":["..."],"safety_notes":["..."],"timeframe":"...","is_code_considerations":["..." ]}}]}}
STRICTLY SEPARATE OBSERVED / PROVIDED CONDITIONS from POSSIBLE DAMAGE MECHANISMS. Observed_conditions must list only what is visible in the image or explicitly stated by the user (e.g., "Concrete spalling with visible reinforcement"). Possible_damage_mechanism.most_likely must be conservatively worded — when evidence is insufficient, write "Cannot be distinguished reliably from the available evidence." or "Possible mechanisms include… (requires field verification)". Never write "The most likely cause is…" unless the user has supplied enough field information to narrow the possibilities; if narrowing, briefly state what evidence justified the narrowing. Use hedged wording such as "May indicate…", "Could be associated with…", "Consistent with…", "Possible…", "Requires field verification…". Never state a cause as definite unless the evidence establishes it. user_measurements must contain only values the user actually provided; if a field is "Not measured", "Not visible", or "Unknown", either omit it or write "<field>: Not provided.". If a photograph does not allow a measurement, write "Not reliably measurable from the submitted image." Never fabricate crack widths, crack depths, concrete strength, reinforcement size, hidden reinforcement condition, structural capacity, load capacity, residual capacity, or safety percentages. Do NOT emit numerical confidence values in user-facing text; evidence_confidence.percentage is optional and not displayed to end users. preliminary_severity is a PRELIMINARY OBSERVED DAMAGE SEVERITY only — not a final safety classification; use "Indeterminate" when evidence is insufficient. evidence_quality rates image quality, visibility, number of photographs, and completeness of provided information. missing_information must clearly list additional info that would improve the assessment (e.g., crack width, member dimensions, drawings, concrete strength, reinforcement details, loading history, damage extent). NEVER recommend demolition on the basis of image-only evidence and NEVER definitively declare a structure safe or unsafe from a photograph alone.
Internally classify every important statement as OBSERVED, ESTIMATED, INFERRED, or UNKNOWN, then communicate those labels naturally. Evidence_detected may include only visible image features or explicit user fields. If the image is poor, irrelevant, ambiguous, or no structural evidence is available, set assessment_status and damage_classification to "Insufficient evidence", preliminary_severity to "Indeterminate", evidence_quality to "Poor" or "Insufficient", lower confidence, explain what is visible/provided and what is unknown, request better photographs (overall context, close-up, scale, multiple angles), and do not issue a confident diagnosis. Any reported incident (blast, fire, earthquake, impact, etc.) is user-provided context, not visual proof of that damage type.
Provide 3 concise ranked pathways only when evidence supports decision support; otherwise keep recommendations conservative and verification-led. Connect each pathway to evidence, state when it is appropriate, and when it must not be used. For each option identify applicable Indian Standard families only when relevant (for example IS 456 for reinforced concrete, IS 15988 for evaluation/strengthening, IS 13920 for ductile detailing, IS 800 for structural steel, IS 1893 for seismic, IS 1905 for masonry, or IS 516 for material testing). Do not invent clause numbers; say current editions and project-specific provisions must be verified.
Assessment: {payload.model_dump_json(exclude={"image_base64"})}"""
    chat = LlmChat(api_key=key, session_id=str(uuid.uuid4()), system_message="You provide cautious, evidence-aware structural rehabilitation decision support.").with_model("openai", "gpt-5.4-mini")
    contents = []
    if payload.image_base64 and payload.image_mime in {"image/jpeg", "image/png", "image/webp"}:
        from emergentintegrations.llm.chat import ImageContent
        contents.append(ImageContent(image_base64=payload.image_base64.split(",")[-1]))
        prompt += "\nA site photo is attached. Use it only as supplementary visual evidence; do not infer hidden damage."
    output = []
    async for event in chat.stream_message(UserMessage(text=prompt, file_contents=contents)):
        if isinstance(event, TextDelta):
            output.append(event.content)
    try:
        return clean_json_text("".join(output))
    except Exception as exc:
        logging.exception("AI returned invalid JSON")
        raise HTTPException(status_code=502, detail="The AI returned an unreadable assessment. Please try again.") from exc


@api_router.post("/assessments", response_model=Assessment)
async def create_assessment(payload: AssessmentCreate):
    advice = await generate_advice(payload)
    normalized = normalize_advice(advice, payload.component_type)
    record = {
        "id": str(uuid.uuid4()), "structure_name": payload.structure_name,
        "component_type": payload.component_type, "integrity_status": payload.integrity_status,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "recommendations": normalized["recommendations"],
        "immediate_actions": normalized["immediate_actions"],
        "disclaimer": "AI provides preliminary decision support, not structural certification. A licensed structural engineer must inspect, verify calculations, assess load paths, and approve all work; this report does not guarantee safety, occupancy, capacity, or repair performance.",
        "assessment_status": normalized["assessment_status"],
        "damage_classification": normalized["damage_classification"],
        "quick_assessment": normalized["quick_assessment"],
        "evidence_confidence": normalized["evidence_confidence"],
        "possible_damage_mechanism": normalized["possible_damage_mechanism"],
        "limitations": normalized["limitations"],
        "required_verification": normalized["required_verification"],
        "damage_type": payload.damage_type,
        "construction_material": payload.construction_material,
        "observed_conditions": normalized["observed_conditions"],
        "user_measurements": normalized["user_measurements"],
        "preliminary_severity": normalized["preliminary_severity"],
        "evidence_quality": normalized["evidence_quality"],
        "evidence_label": normalized["evidence_label"],
        "missing_information": normalized["missing_information"],
        "recovery_pathway": normalized["recovery_pathway"],
        "priority_level": normalized["priority_level"],
        "recovery_pathway_reason": normalized["recovery_pathway_reason"],
        "workflow_steps": normalized["workflow_steps"],
        "decision_table": normalized["decision_table"],
    }
    await db.assessments.insert_one(record.copy())
    return record

@api_router.get("/assessments", response_model=List[Assessment])
async def list_assessments():
    records = await db.assessments.find({}, {"_id": 0}).sort("created_at", -1).to_list(50)
    return records

async def get_assessment_or_404(assessment_id: str) -> dict:
    record = await db.assessments.find_one({"id": assessment_id}, {"_id": 0})
    if not record:
        raise HTTPException(status_code=404, detail="Assessment not found")
    return record

def report_filename(record: dict, extension: str) -> str:
    safe_name = "".join(c if c.isalnum() or c in "-_" else "_" for c in record["structure_name"])[:48]
    return f"{safe_name}_resructai_report.{extension}"

WATERMARK_PRIMARY = "RESTRUCT AI"
WATERMARK_SECONDARY = "AI-POWERED STRUCTURAL DAMAGE ASSESSMENT"

def _pdf_watermark(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(Color(0.75, 0.75, 0.78, alpha=0.18))
    canvas.setFont("Helvetica-Bold", 72)
    canvas.translate(A4[0] / 2, A4[1] / 2)
    canvas.rotate(35)
    canvas.drawCentredString(0, 0, WATERMARK_PRIMARY)
    canvas.setFont("Helvetica", 14)
    canvas.setFillColor(Color(0.75, 0.75, 0.78, alpha=0.22))
    canvas.drawCentredString(0, -46, WATERMARK_SECONDARY)
    canvas.restoreState()
    canvas.saveState()
    canvas.setFillColor(Color(0.4, 0.4, 0.45, alpha=0.7))
    canvas.setFont("Helvetica", 8)
    canvas.drawString(42, 20, f"{WATERMARK_PRIMARY} · {WATERMARK_SECONDARY}")
    canvas.drawRightString(A4[0] - 42, 20, f"Page {doc.page}")
    canvas.restoreState()

@api_router.get("/assessments/{assessment_id}/report/{file_format}")
async def download_report(assessment_id: str, file_format: str):
    record = await get_assessment_or_404(assessment_id)
    if file_format not in {"pdf", "docx", "xlsx"}:
        raise HTTPException(status_code=400, detail="Choose pdf, docx, or xlsx")
    buffer = BytesIO()
    title = f"ResructAI | {record['structure_name']}"
    if file_format == "pdf":
        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
        quick = record.get("quick_assessment", {}); confidence = record.get("evidence_confidence", {}); mechanism = record.get("possible_damage_mechanism", {})
        story = [Paragraph("RESRUCTAI / STRUCTURAL REVIEW", styles["Title"]), Paragraph(title, styles["Heading2"]), Paragraph(f"{record['component_type']} · {record['integrity_status']} · {record['created_at'][:10]}", styles["Normal"]), Spacer(1, 16), Paragraph("ASSESSMENT OVERVIEW", styles["Heading3"]), Paragraph(f"Damage type: {record.get('damage_type', 'Unknown')} | Material: {record.get('construction_material', 'Unknown')} | Preliminary observed severity: {record.get('preliminary_severity', 'Indeterminate')} | Evidence quality: {record.get('evidence_quality', 'Insufficient')} ({record.get('evidence_label', 'Insufficient evidence')})", styles["BodyText"]), Paragraph("QUICK ASSESSMENT", styles["Heading3"]), Paragraph(f"Element: {quick.get('element', record['component_type'])} | Condition: {quick.get('apparent_condition', 'Insufficient evidence')} | Severity: {quick.get('severity', record.get('damage_classification', 'Insufficient evidence'))} | Evidence: {record.get('evidence_label', 'Insufficient evidence')}", styles["BodyText"]), Paragraph("PRELIMINARY RECOVERY PATHWAY", styles["Heading3"]), Paragraph(f"Pathway: {record.get('recovery_pathway', 'Further Investigation Required')} | Priority: {record.get('priority_level', 'Indeterminate')}", styles["BodyText"]), Paragraph("Why this pathway was selected: " + record.get("recovery_pathway_reason", ""), styles["BodyText"]), Paragraph("OBSERVED / PROVIDED CONDITIONS", styles["Heading3"]), Paragraph("; ".join(record.get("observed_conditions", [])) or "No observed conditions recorded from image or user input.", styles["BodyText"]), Paragraph("USER-PROVIDED MEASUREMENTS", styles["Heading3"]), Paragraph("; ".join(record.get("user_measurements", [])) or "No measurements provided by the user.", styles["BodyText"]), Paragraph("EVIDENCE (OBSERVED / INFERRED / UNKNOWN)", styles["Heading3"]), Paragraph("Observed / provided: " + "; ".join(confidence.get("evidence_detected", [])), styles["BodyText"]), Paragraph("Estimated / inferred: " + "; ".join(confidence.get("estimated_inferred", [])), styles["BodyText"]), Paragraph("Cannot be determined: " + "; ".join(confidence.get("cannot_be_determined", [])), styles["BodyText"]), Paragraph("POSSIBLE DAMAGE MECHANISMS", styles["Heading3"]), Paragraph(f"Possible mechanisms: {mechanism.get('most_likely', 'Cannot be distinguished reliably from the available evidence.')}", styles["BodyText"]), Paragraph("Alternatives: " + ("; ".join(mechanism.get('alternatives', [])) or "None established from available evidence."), styles["BodyText"]), Paragraph("MISSING INFORMATION", styles["Heading3"]), Paragraph("; ".join(record.get("missing_information", [])), styles["BodyText"]), Paragraph("IMMEDIATE ACTIONS", styles["Heading3"]), ListFlowable([ListItem(Paragraph(x, styles["BodyText"])) for x in record.get("immediate_actions", [])], bulletType="bullet"), Spacer(1, 8), Paragraph("POSSIBLE TECHNIQUES — SUBJECT TO ENGINEERING VERIFICATION", styles["Heading3"])]
        for item in record.get("recommendations", []):
            story += [Spacer(1, 12), Paragraph(f"{item.get('rank')}. {item.get('technique', '')} — {item.get('urgency', '')}", styles["Heading3"]), Paragraph(item.get("reasoning", ""), styles["BodyText"]), Paragraph("Evidence link: " + item.get("evidence_link", ""), styles["BodyText"]), Paragraph("Appropriate when: " + item.get("appropriate_when", ""), styles["BodyText"]), Paragraph("Do not use when: " + item.get("do_not_use_when", ""), styles["BodyText"]), Paragraph("Materials: " + "; ".join(item.get("materials", [])), styles["BodyText"]), Paragraph("Procedure: " + "; ".join(item.get("procedure", [])), styles["BodyText"]), Paragraph("Safety: " + "; ".join(item.get("safety_notes", [])), styles["BodyText"]), Paragraph("IS code considerations: " + "; ".join(item.get("is_code_considerations", [])), styles["BodyText"])]
        story += [Spacer(1, 16), Paragraph("STEP-BY-STEP RECOVERY WORKFLOW", styles["Heading3"]), ListFlowable([ListItem(Paragraph(f"Step {i+1}. {s}", styles["BodyText"])) for i, s in enumerate(record.get("workflow_steps", []))], bulletType="bullet"), Spacer(1, 8), Paragraph("DECISION REFERENCE TABLE", styles["Heading3"])]
        for row in record.get("decision_table", []):
            story += [Paragraph(f"Evidence: {row.get('evidence','')} → Pathway: {row.get('pathway','')} → Verification: {row.get('verification','')}", styles["BodyText"])]
        story += [Spacer(1, 12), Paragraph("LIMITATIONS OF IMAGE-BASED ASSESSMENT", styles["Heading3"]), Paragraph("; ".join(record.get("limitations", [])), styles["BodyText"]), Paragraph("A photograph alone generally cannot establish: hidden reinforcement condition, crack depth, actual structural capacity, material strength, load capacity, internal damage, foundation condition, complete load path, long-term durability, or actual safety/occupancy status without supporting engineering analysis.", styles["BodyText"]), Paragraph("REQUIRED VERIFICATION", styles["Heading3"]), Paragraph("; ".join(record.get("required_verification", [])), styles["BodyText"]), Paragraph("PROFESSIONAL REVIEW REQUIRED", styles["Heading3"]), Paragraph("Replacement or demolition decisions require detailed structural assessment, consideration of safety, feasibility, economics, regulatory requirements and approval by the appropriate qualified professionals. This report is preliminary decision support only and is not an engineering certification.", styles["BodyText"]), Paragraph(record["disclaimer"], styles["BodyText"])]
        doc.build(story, onFirstPage=_pdf_watermark, onLaterPages=_pdf_watermark)
        media_type = "application/pdf"
    elif file_format == "docx":
        document = Document()
        header = document.sections[0].header.paragraphs[0]
        header.text = ""
        run = header.add_run(f"{WATERMARK_PRIMARY}  ·  {WATERMARK_SECONDARY}")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xA0, 0xA0, 0xA8)
        footer = document.sections[0].footer.paragraphs[0]
        footer.text = ""
        frun = footer.add_run(f"{WATERMARK_PRIMARY} — Preliminary Engineering Decision Support · Not a Structural Certification")
        frun.italic = True
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(0xA0, 0xA0, 0xA8)
        document.add_heading("RESRUCTAI / STRUCTURAL REVIEW", 0); document.add_heading(title, 1)
        document.add_paragraph(f"{record['component_type']} · {record['integrity_status']} · {record['created_at'][:10]}")
        quick = record.get("quick_assessment", {}); confidence = record.get("evidence_confidence", {}); mechanism = record.get("possible_damage_mechanism", {})
        document.add_heading("Assessment overview", 2); document.add_paragraph(f"Damage type: {record.get('damage_type', 'Unknown')} | Material: {record.get('construction_material', 'Unknown')} | Preliminary observed severity: {record.get('preliminary_severity', 'Indeterminate')} | Evidence quality: {record.get('evidence_quality', 'Insufficient')} ({record.get('evidence_label', 'Insufficient evidence')})")
        document.add_heading("Quick assessment", 2); document.add_paragraph(f"Element: {quick.get('element', record['component_type'])} | Condition: {quick.get('apparent_condition', 'Insufficient evidence')} | Severity: {record.get('damage_classification', 'Insufficient evidence')} | Evidence: {record.get('evidence_label', 'Insufficient evidence')}")
        document.add_heading("Preliminary recovery pathway", 2); document.add_paragraph(f"Pathway: {record.get('recovery_pathway', 'Further Investigation Required')} | Priority: {record.get('priority_level', 'Indeterminate')}"); document.add_paragraph("Why this pathway was selected: " + record.get("recovery_pathway_reason", ""))
        document.add_heading("Observed / provided conditions", 2); document.add_paragraph("; ".join(record.get("observed_conditions", [])) or "No observed conditions recorded from image or user input.")
        document.add_heading("User-provided measurements", 2); document.add_paragraph("; ".join(record.get("user_measurements", [])) or "No measurements provided by the user.")
        document.add_heading("Evidence (observed / inferred / unknown)", 2); document.add_paragraph("Observed / provided: " + "; ".join(confidence.get("evidence_detected", []))); document.add_paragraph("Estimated / inferred: " + "; ".join(confidence.get("estimated_inferred", []))); document.add_paragraph("Cannot be determined: " + "; ".join(confidence.get("cannot_be_determined", [])))
        document.add_heading("Possible damage mechanisms", 2); document.add_paragraph("Possible mechanisms: " + mechanism.get('most_likely', 'Cannot be distinguished reliably from the available evidence.')); document.add_paragraph("Alternatives: " + ("; ".join(mechanism.get('alternatives', [])) or "None established from available evidence."))
        document.add_heading("Missing information", 2); document.add_paragraph("; ".join(record.get("missing_information", [])))
        document.add_heading("Immediate actions", 2)
        for action in record.get("immediate_actions", []): document.add_paragraph(action, style="List Bullet")
        document.add_heading("Possible techniques — subject to engineering verification", 2)
        for item in record.get("recommendations", []):
            document.add_heading(f"{item.get('rank')}. {item.get('technique', '')} — {item.get('urgency', '')}", 2)
            document.add_paragraph(item.get("reasoning", "")); document.add_paragraph("Evidence link: " + item.get("evidence_link", "")); document.add_paragraph("Appropriate when: " + item.get("appropriate_when", "")); document.add_paragraph("Do not use when: " + item.get("do_not_use_when", "")); document.add_paragraph("Materials: " + "; ".join(item.get("materials", []))); document.add_paragraph("Procedure: " + "; ".join(item.get("procedure", []))); document.add_paragraph("Safety: " + "; ".join(item.get("safety_notes", []))); document.add_paragraph("IS code considerations: " + "; ".join(item.get("is_code_considerations", [])))
        document.add_heading("Step-by-step recovery workflow", 2)
        for i, s in enumerate(record.get("workflow_steps", [])): document.add_paragraph(f"Step {i+1}. {s}", style="List Number")
        document.add_heading("Decision reference table", 2)
        for row in record.get("decision_table", []): document.add_paragraph(f"Evidence: {row.get('evidence','')} → Pathway: {row.get('pathway','')} → Verification: {row.get('verification','')}")
        document.add_heading("Limitations of image-based assessment", 2); document.add_paragraph("; ".join(record.get("limitations", []))); document.add_paragraph("A photograph alone generally cannot establish hidden reinforcement condition, crack depth, structural capacity, material strength, load capacity, internal damage, foundation condition, complete load path, long-term durability, or actual safety/occupancy status without supporting engineering analysis.")
        document.add_heading("Required verification", 2); document.add_paragraph("; ".join(record.get("required_verification", []))); document.add_heading("Professional review required", 2); document.add_paragraph("Replacement or demolition decisions require detailed structural assessment, consideration of safety, feasibility, economics, regulatory requirements and approval by the appropriate qualified professionals. This report is preliminary decision support only and is not an engineering certification."); document.add_paragraph(record["disclaimer"]); document.save(buffer)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Assessment"
        sheet.oddHeader.center.text = f"&\"Arial,Bold\"&12&K808080{WATERMARK_PRIMARY}\n&\"Arial,Italic\"&9&K909098{WATERMARK_SECONDARY}"
        sheet.oddFooter.center.text = f"&\"Arial,Italic\"&8&K909098{WATERMARK_PRIMARY} — Preliminary Engineering Decision Support · Not a Structural Certification"
        sheet.oddFooter.right.text = "&\"Arial\"&8&K909098Page &P of &N"
        sheet.append(["ResructAI structural review", record["structure_name"]]); sheet.append(["Damage type", record.get("damage_type", "Unknown")]); sheet.append(["Component", record["component_type"]]); sheet.append(["Material", record.get("construction_material", "Unknown")]); sheet.append(["Integrity", record["integrity_status"]]); sheet.append([])
        sheet.append(["Assessment status", record.get("assessment_status", "Insufficient evidence")]); sheet.append(["Damage classification", record.get("damage_classification", "Insufficient evidence")]); sheet.append(["Preliminary observed severity", record.get("preliminary_severity", "Indeterminate")]); sheet.append(["Evidence quality", record.get("evidence_quality", "Insufficient")]); sheet.append(["Evidence label", record.get("evidence_label", "Insufficient evidence")]); sheet.append(["Preliminary recovery pathway", record.get("recovery_pathway", "Further Investigation Required")]); sheet.append(["Priority level", record.get("priority_level", "Indeterminate")]); sheet.append(["Why this pathway", record.get("recovery_pathway_reason", "")]); confidence = record.get("evidence_confidence", {}); sheet.append(["Observed / provided conditions", "\n".join(record.get("observed_conditions", []))]); sheet.append(["User-provided measurements", "\n".join(record.get("user_measurements", []))]); sheet.append(["Evidence detected", "\n".join(confidence.get("evidence_detected", []))]); sheet.append(["Estimated / inferred", "\n".join(confidence.get("estimated_inferred", []))]); sheet.append(["Cannot be determined", "\n".join(confidence.get("cannot_be_determined", []))]); sheet.append(["Missing information", "\n".join(record.get("missing_information", []))]); sheet.append([])
        sheet.append(["Rank", "Technique (subject to engineering verification)", "Urgency", "Reasoning", "Evidence link", "Appropriate when", "Do not use when", "Materials", "Procedure", "Safety notes", "IS code considerations"])
        for item in record.get("recommendations", []): sheet.append([item.get("rank"), item.get("technique"), item.get("urgency"), item.get("reasoning"), item.get("evidence_link"), item.get("appropriate_when"), item.get("do_not_use_when"), "\n".join(item.get("materials", [])), "\n".join(item.get("procedure", [])), "\n".join(item.get("safety_notes", [])), "\n".join(item.get("is_code_considerations", []))])
        sheet.append([]); sheet.append(["Step-by-step recovery workflow"])
        for i, s in enumerate(record.get("workflow_steps", [])): sheet.append([f"Step {i+1}", s])
        sheet.append([]); sheet.append(["Decision reference table"]); sheet.append(["Evidence condition", "Preliminary pathway", "Required verification"])
        for row in record.get("decision_table", []): sheet.append([row.get("evidence", ""), row.get("pathway", ""), row.get("verification", "")])
        sheet.append([]); sheet.append(["Required verification", "\n".join(record.get("required_verification", []))]); sheet.append(["Limitations", "\n".join(record.get("limitations", []))]); sheet.append(["Disclaimer", record["disclaimer"]]); sheet.column_dimensions["A"].width = 30; sheet.column_dimensions["B"].width = 50
        workbook.save(buffer); media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    buffer.seek(0)
    return StreamingResponse(buffer, media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{report_filename(record, file_format)}"'})

@api_router.post("/status", response_model=StatusCheck)
async def create_status_check(input: StatusCheckCreate):
    status_dict = input.model_dump()
    status_obj = StatusCheck(**status_dict)
    
    # Convert to dict and serialize datetime to ISO string for MongoDB
    doc = status_obj.model_dump()
    doc['timestamp'] = doc['timestamp'].isoformat()
    
    _ = await db.status_checks.insert_one(doc)
    return status_obj

@api_router.get("/status", response_model=List[StatusCheck])
async def get_status_checks():
    # Exclude MongoDB's _id field from the query results
    status_checks = await db.status_checks.find({}, {"_id": 0}).to_list(1000)
    
    # Convert ISO string timestamps back to datetime objects
    for check in status_checks:
        if isinstance(check['timestamp'], str):
            check['timestamp'] = datetime.fromisoformat(check['timestamp'])
    
    return status_checks

# Include the router in the main app
app.include_router(api_router)

cors_origins = [
    origin.strip()
    for origin in os.getenv("CORS_ORIGINS", "").split(",")
    if origin.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
